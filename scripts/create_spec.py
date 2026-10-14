import json,re
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
inventory=json.loads((root/'reference/inventory.json').read_text(encoding='utf8'))
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 2'].font.size=Pt(12)
def para(text,bold=False):
 p=doc.add_paragraph();p.add_run(text).bold=bold;return p
def heading(text):doc.add_heading(text,1)
def page_data(p):return json.loads((root/'reference/pages'/('home.json' if p['route']=='/' else p['route'][1:].replace('/','__')+'.json')).read_text(encoding='utf8'))
def table(rows):
 t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Table Grid'
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for c,text in zip(cells,row):
   c.text=str(text)
   for run in c.paragraphs[0].runs:run.font.size=Pt(9);run.bold=i==0
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EEF3');c._tc.get_or_add_tcPr().append(shade)
  trPr=t.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
 return t

doc.add_paragraph('Winbuzz website replica specification','Title')
para('Page inventory and Next.js implementation reference',True)
para('Reference website: https://officialwinbuzz.com/\nInspection date: 13 October 2026\nPrepared for the website development assignment')
para(f'This document records the public page structure, content, visual system, navigation, images, and interaction requirements for recreating the reference website. The local implementation contains {len(inventory["pages"])} captured routes and uses the original downloaded image files. External links and form submissions share one configurable redirect destination.')
heading('Scope and source files')
para('The public WordPress site is reproduced through the Next.js App Router. Each original route is backed by a local content snapshot. The page and post sitemaps identify the main pages and articles; category, tag, author, and linked pagination routes complete the public archive inventory. RSS feeds, XML sitemaps, administration, and API endpoints are infrastructure rather than visual pages.')
para('Exact page copy, markup, image placement, and original CSS are preserved in reference/pages/*.json. The machine-readable reference/inventory.json contains route titles, descriptions, section headings, original link destinations, image alt text, form fields, and the source-to-local asset mapping. These files provide the implementation details alongside this document.')
heading('Visual design and language')
table([['Token','Original CSS value','Usage'],['Background','oklch(16% .02 270)','Dark navy page canvas'],['Foreground','oklch(97% .01 90)','Warm near-white headings and text'],['Card','oklch(21% .025 270)','Login panels and content cards'],['Primary','oklch(82% .15 85)','Gold accents and calls to action'],['Muted foreground','oklch(72% .02 90)','Secondary paragraphs'],['Border','oklch(32% .03 270)','Subtle navy borders'],['Radius','.75rem','Default rounded corners']])
para('Typography declares Inter, system-ui, sans-serif. The source does not supply an Inter font file in the captured styles, so the system fallback is used if Inter is unavailable. Preserve the declared font stack rather than substitute a different display font. Body copy uses regular weight; navigation and buttons use medium or semibold; headings use bold. Text sizes include 12, 14, 16, 18, 20, 24, 30, 36, and 48 pixels. The home hero is 30 pixels on small screens, 36 on medium screens, and 48 on large screens; its line height is 1.05. On mobile, the theme reduces paragraph font size to 14 pixels while retaining the utility line heights.')
para('All inspected page copy and navigation are English, with en-US declared on the document. References to Hindi and other Indian languages describe platform support; there is no visible language selector on this public site. The currency mentioned in copy is Indian Rupees.')
para('Gold buttons use a pale-yellow to deeper-gold gradient, dark text, rounded or pill corners, and a soft gold glow. Outlined buttons have a thin gold border. Cards use a navy gradient or dark surface, light borders, and rounded corners. Images use cricket stadium, casino, mobile app, and payment imagery. WhatsApp green and Telegram blue are reserved for floating support buttons.')
heading('Layout and responsive behavior')
para('The header has a left logo, central desktop navigation, and right signup and login buttons. Desktop navigation contains Home, About Us, Sports, Winbuzz App, Blog, and Faq. At smaller widths it becomes a slide-out menu with a hamburger, close button, and overlay. The theme uses a 1200-pixel desktop navigation breakpoint. Content uses centered containers up to 1280 pixels, horizontal padding of 16 pixels, and section spacing commonly 64 pixels on small screens and 96 pixels from the medium breakpoint.')
para('The home hero combines a full-width cricket background with dark overlays, a left heading and introduction, gold ID button, three assurance labels, and a right login card. On small screens the content and login card stack. Subsequent sections alternate image-and-copy blocks, multi-column card grids, tables, lists, testimonials, and FAQs. Medium grids use two or three columns, then collapse to one column on small screens. Header height and viewport height are measured for the menu and sticky header. Hash navigation uses a 100-pixel scroll offset.')
heading('Navigation and interactions')
table([['Control','Required behavior'],['Home','Navigate to /'],['About Us','Navigate to /about-us'],['Sports','Navigate to /#sports'],['Winbuzz App','Navigate to /#app'],['Blog','Navigate to /blog and its pagination'],['Faq','Navigate to /#faq'],['Footer main links','Navigate to their matching local routes'],['External CTA or support link','Navigate through /out to EXTERNAL_REDIRECT_URL'],['Form submission','Discard all values and redirect through /out'],['FAQ','Open first item initially; opening another closes the previous item'],['Mobile menu','Toggle menu; close on overlay, menu link, or Escape'],['Reviews','Switch between the original reviews through pagination and next or previous buttons'],['Score tabs','Local Recent, Live, Upcoming panel; provider link uses /out']])
para('Original outgoing destinations include WhatsApp through walink.ai, the Telegram Manualid bot, winbuzz.app, winbuzzlive.live, Trustpilot, and the Crictimes score provider. Every external navigation link in the replica uses the same server-side environment variable EXTERNAL_REDIRECT_URL. The default example is https://example.com. The /out route accepts GET and POST and returns a 303 redirect. Submitted usernames, passwords, email addresses, phone numbers, and other form values are never forwarded, persisted, or appended to the destination URL. No OTP, real login, account registration, KYC, wallet, payment, or gambling transaction is implemented.')
para('The score iframe is rendered locally at /scores using the original provider CSS and downloaded Roboto and Lato fonts. Recent, Live, and Upcoming tabs display real match cards from /api/scores, refreshed every 30 seconds. Saved provider data supplies a fallback when the feed is unavailable. Every match-card and provider link uses /out. The review carousel uses the original downloaded Swiper library and original container breakpoints, 400 millisecond transitions, 20 pixel spacing, 3 second autoplay, drag and swipe, navigation arrows, and pagination. Autoplay pauses on hover and respects reduced motion. Analytics and trackers are removed. The locally generated utility stylesheet replaces the runtime Tailwind CDN script.')
doc.add_page_break();heading('Main page inventory')
core=['/','/about-us','/login','/sign-up','/contact-us','/disclaimer','/privacy-policy','/responsible-gaming-policy','/rules-and-regulations','/terms-and-conditions','/blog']
for route in core:
 p=next((p for p in inventory['pages'] if p['route']==route),None)
 if not p:continue
 doc.add_heading(('Home' if route=='/' else route[1:].replace('-',' ').title()),2)
 para(route+' | '+p['title'],True)
 if p.get('description'):para(p['description'])
 sections=p.get('sections',p['headings']);sections=[h['text'] for h in sections if h['level']<= (2 if route=='/' else 3) and h['text'] not in ['Quick Links','Winbuzz Platform']]
 if route=='/' and p.get('blocks'):sections=[b['heading']+(' (#'+b['id']+')' if b['id'] else '') for b in p['blocks'] if b['heading']]
 para('Content in order: '+ '; '.join(dict.fromkeys(sections)))
 if p.get('fields'):para('Visible fields: '+', '.join(f["label"]+' ('+f['type']+')' for f in p['fields']))
 if route=='/blog':para('Article cards show feature image, title, excerpt, and a detail link. Numbered pagination opens /blog/page/2 through /blog/page/7. Article routes use their original slugs.')

heading('Shared footer and support controls')
para('The footer includes Winbuzz Platform outgoing links, a Quick Links heading, logo, introductory paragraph, Trustpilot link, and three columns of links and labels. Local links include Login, Signup, Blog, About Us, Contact Us, Disclaimer, Privacy Policy, Responsible Gaming, Rules and Regulations, and Terms and Conditions. Some original footer labels are plain text rather than clickable links. The footer ends with the original keyword list and copyright text. Floating WhatsApp and Telegram controls remain visible above content and route to the configured destination.')
heading('Original images and asset storage')
para(f'The capture stores {len(inventory["assets"])} source assets locally. Asset filenames use a stable hash of their original URL plus the file extension. All image elements reference /assets files, and valid CSS background images are also downloaded. Alternative responsive image candidates are removed so the locally downloaded source image is used consistently. The inventory preserves the URL-to-file mapping and each page retains alt text.')
home=next(p for p in inventory['pages'] if p['route']=='/')
table([['Home image','Local asset']]+[[im.get('alt') or 'Logo or decoration',im.get('src','')] for im in home['images'] if im.get('alt')][:13])
doc.add_page_break();heading('Article page inventory')
articles=[p for p in inventory['pages'] if p['route'] not in core and not any(p['route'].startswith(prefix) for prefix in ['/tag/','/category/','/author/','/blog/page/'])]
para(f'{len(articles)} captured article routes use the article template. Each contains its original title, editorial copy, section headings, lists or tables where present, links, and shared header and footer. The following entries list the route and page-specific subject areas; exact body copy is retained in the local page snapshots.')
for p in articles:
 para(p['route'],True)
 title=p['title'].replace(' - Winbuzz','');para(title)
 sections=[h['text'] for h in p.get('sections',p['headings']) if h['level'] in [2,3] and h['text'] not in ['Quick Links','Winbuzz Platform']]
 if sections:para('Sections: '+'; '.join(dict.fromkeys(sections)))
 else:para(p.get('description') or 'Editorial article with the original page title and body text.')
heading('Archive and pagination page inventory')
archives=[p for p in inventory['pages'] if any(p['route'].startswith(prefix) for prefix in ['/tag/','/category/','/author/','/blog/page/'])]
para('Tag archives list articles associated with the named tag; category archives list the category posts; author archives list the author posts. Each uses the original archive title, article summaries, feature images, and pagination where present. Blog pagination continues the blog listing. Every archive route is listed below.')
table([['Route','Contains']]+[[p['route'],p['title'].replace(' - Winbuzz','')] for p in archives])
heading('Implementation and validation requirements')
para('Use npm install, copy .env.example to .env.local, set EXTERNAL_REDIRECT_URL, and run npm run dev. Production commands are npm run build and npm start. The application uses Next.js 16 with React 19 and the App Router. Server rendering reads the captured local page data; React handles menu, FAQ coordination, forms, score tabs, and review navigation. Unknown routes return a local 404 page. No WordPress hosting is required.')
para('Run npm run audit to validate internal link coverage, image locality, external-link routing, tracker removal, and form privacy. With AUDIT_BASE_URL set to the preview origin, the audit additionally checks every route returns HTTP 200, the /out GET and POST responses redirect with status 303, and a missing route returns 404. Review desktop and mobile layouts, menu opening and closing, hash navigation, FAQ expansion, article pagination, form redirects, and local asset loading before assignment submission.')
heading('Responsive browser verification')
para('All 271 captured routes were tested in the browser at widths of 320, 375, 425, 768, 1024, 1280, 1440, and 1920 pixels, for 2168 responsive checks. No horizontal overflow, broken eagerly loaded images, or remote image references were found. Sixteen source layouts were compared at those widths for static text geometry, typography, colors, and corners, with no remaining measured differences after correcting utility corner radii and testimonial spacing. This is a template comparison plus a complete route sweep, rather than a pixel comparison of every full page. Direct viewport comparison with the third-party score provider was blocked by browser permission; its downloaded styles and local implementation were tested separately. Evidence and screenshots are retained in docs/visual-qa.')
heading('Reference sources')
for url in ['https://officialwinbuzz.com/','https://officialwinbuzz.com/sitemap_index.xml','https://officialwinbuzz.com/page-sitemap.xml','https://officialwinbuzz.com/post-sitemap.xml','https://officialwinbuzz.com/category-sitemap.xml','https://officialwinbuzz.com/post_tag-sitemap.xml','https://officialwinbuzz.com/author-sitemap.xml']:para(url)
doc.save(root/'docs/Winbuzz-website-specification.docx')
print('Created specification for',len(inventory['pages']),'routes')
